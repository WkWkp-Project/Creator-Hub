"""Content Asset Suite endpoints — per-campaign access control.

Access model:
  - admin   : sees & manages every campaign; the only role that can create,
              delete, and assign access.
  - manager : sees & edits only campaigns assigned to them (assigned_user_ids).
  - viewer  : sees (read-only) only campaigns assigned to them.
"""
import io
import json
import re
from datetime import datetime

import openpyxl
import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from rapidfuzz import fuzz
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from .. import models
from ..config import get_settings
from ..campaign_report import render_campaign_report
from ..content_asset_models import (
    DEFAULT_INPUT_FILES,
    ContentAsset,
    ContentAssetCreate,
    ContentAssetList,
    ContentAssetOut,
    ContentAssetSummaryList,
    ContentAssetUpdate,
)
from ..database import get_db
from ..deps import get_current_user, require_admin
from ..directory_models import Brand
from ..file_validation import validate_spreadsheet_upload

router = APIRouter(prefix="/api/assets", tags=["content-asset"])
settings = get_settings()


def _enforce_company(db: Session, obj: ContentAsset) -> None:
    """A campaign's Company is governed by its Brand: if the campaign is under a
    brand that has a company, force client_name to that company so Company →
    Brand → Campaign stays consistent."""
    if obj.brand_id:
        br = db.get(Brand, obj.brand_id)
        if br and br.company:
            obj.client_name = br.company

# Governance fields only an admin may change (a manager editing a campaign must
# not be able to grant access or reassign ownership).
ADMIN_ONLY_FIELDS = {"assigned_user_ids", "owner_email", "owner_name", "responsible_member_id", "responsible_member_ids"}


def _can_read(user: models.User, asset: ContentAsset) -> bool:
    return user.role == "admin" or user.id in (asset.assigned_user_ids or [])


def _can_edit(user: models.User, asset: ContentAsset) -> bool:
    return user.role == "admin" or (user.role == "manager" and user.id in (asset.assigned_user_ids or []))


_BUDGET_KEYS = (
    # Legacy Section B fields.
    "rate", "gen_code_price", "boosting_cost",
    # Section B v2 "KOLs Confirmed" fields.
    "kol_price", "gencode_boosting", "cart_added", "buy_asset", "outside_shooting",
)


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _safe_url(u) -> str:
    """Allow only http(s) or root-relative URLs; strip javascript:/data:/vbscript:
    so a stored link can never become an XSS sink when rendered as href."""
    s = str(u or "").strip()
    if not s:
        return ""
    return s if s.lower().startswith(("http://", "https://", "/")) else ""


def _sanitize_urls(obj: ContentAsset) -> None:
    """Neutralise any unsafe URLs stored on a campaign (drive folder + KOL links)."""
    obj.drive_folder_url = _safe_url(obj.drive_folder_url)
    if obj.kols:
        cleaned = []
        for k in obj.kols:
            if not isinstance(k, dict):
                cleaned.append(k)
                continue
            row = dict(k)
            if "link" in row:
                row["link"] = _safe_url(row.get("link"))
            if "profile_link" in row:
                row["profile_link"] = _safe_url(row.get("profile_link"))
            if isinstance(row.get("links"), dict):
                row["links"] = {p: _safe_url(u) for p, u in row["links"].items() if _safe_url(u)}
            cleaned.append(row)
        obj.kols = cleaned
    if obj.input_files:
        obj.input_files = [{**f, "drive_url": _safe_url(f.get("drive_url"))} if "drive_url" in f else f for f in obj.input_files]


def _log(db: Session, asset_id: int, user: models.User, action: str, summary: str = "", detail: dict | None = None) -> None:
    """Append a campaign audit-trail entry (via the central recorder)."""
    from .. import audit
    audit.record(db, entity="campaign", entity_id=asset_id, user=user,
                 action=action, summary=summary, detail=detail)


def _redact_budget(asset: ContentAsset, user: models.User) -> None:
    """Enforce budget_show server-side: blank the budget figures a campaign hid
    from the customer (viewer). Mutates the in-memory object only — this runs in
    read paths that never commit, so it does not persist."""
    if user.role != "viewer":
        return
    bshow = asset.budget_show or {}
    # Hiding the grand total must also blank every component for the viewer,
    # otherwise they just sum the rows themselves.
    if bshow.get("total") is False:
        hidden = list(_BUDGET_KEYS)
    else:
        hidden = [k for k in _BUDGET_KEYS if bshow.get(k) is False]
    if hidden:
        asset.kols = [{**k, **{f: "" for f in hidden}} for k in (asset.kols or [])]


def _hydrate_kol_names(asset: ContentAsset, db: Session) -> None:
    """Attach KOL names to campaign rows without granting Directory profile access."""
    ids: set[int] = set()
    for k in asset.kols or []:
        if not isinstance(k, dict) or k.get("name"):
            continue
        try:
            iid = int(k.get("influencer_id") or 0)
        except (TypeError, ValueError):
            continue
        if iid:
            ids.add(iid)
    if not ids:
        return
    names = dict(db.execute(select(models.Influencer.id, models.Influencer.name).where(models.Influencer.id.in_(ids))).all())
    if not names:
        return
    hydrated = []
    for k in asset.kols or []:
        if not isinstance(k, dict) or k.get("name"):
            hydrated.append(k)
            continue
        try:
            iid = int(k.get("influencer_id") or 0)
        except (TypeError, ValueError):
            hydrated.append(k)
            continue
        hydrated.append({**k, "name": names.get(iid, k.get("name", ""))} if iid in names else k)
    asset.kols = hydrated


def _filter_asset_query(stmt, search=None, owner_email=None, status=None, brand_id=None):
    if search:
        like = f"%{search}%"
        stmt = stmt.where(or_(ContentAsset.campaign_name.ilike(like), ContentAsset.client_name.ilike(like)))
    if owner_email:
        stmt = stmt.where(ContentAsset.owner_email == owner_email)
    if brand_id is not None:
        stmt = stmt.where(ContentAsset.brand_id == brand_id)
    if status and status.lower() not in {"all", ""}:
        stmt = stmt.where(ContentAsset.status == status.lower())
    return stmt


def _summary_row(row, user: models.User) -> dict:
    kols = row.kols or []
    approved = posted = pending = overdue = soon = 0
    today = datetime.now().date()
    closed = str(row.status or "").lower() in {"completed", "complete", "done", "success", "succeeded", "cancelled"}
    budget_show = row.budget_show or {}
    hidden_budget = set()
    if user.role == "viewer":
        hidden_budget = (set(_BUDGET_KEYS) if budget_show.get("total") is False
                         else {key for key in _BUDGET_KEYS if budget_show.get(key) is False})
    budget = 0.0
    for kol in kols:
        if not isinstance(kol, dict):
            continue
        state = kol.get("client_approved")
        approved += state == "Approve"
        posted += state == "Posted"
        pending += state == "Pending"
        for key in ("rate", "gen_code_price", "boosting_cost"):
            if key not in hidden_budget:
                budget += _num(kol.get(key))
        if closed or state == "Approve":
            continue
        value = kol.get("post_date") or kol.get("period_to")
        if not value or not re.match(r"^\d{4}-\d{2}-\d{2}", str(value)):
            continue
        try:
            days = (datetime.fromisoformat(str(value)[:10]).date() - today).days
        except ValueError:
            continue
        if days < 0:
            overdue += 1
        elif days <= 7:
            soon += 1

    files = row.input_files or []
    file_count = len(DEFAULT_INPUT_FILES)
    linked = sum(1 for item in files[:file_count] if isinstance(item, dict) and item.get("linked"))
    return {
        "id": row.id,
        "client_name": row.client_name or "",
        "campaign_name": row.campaign_name,
        "brand_id": row.brand_id,
        "responsible_member_ids": row.responsible_member_ids or (
            [row.responsible_member_id] if row.responsible_member_id else []
        ),
        "period_start": row.period_start or "",
        "period_end": row.period_end or "",
        "status": row.status or "draft",
        "can_edit": user.role in {"admin", "manager"},
        "kol_count": len(kols),
        "approved_count": approved,
        "posted_count": posted,
        "pending_count": pending,
        "overdue_count": overdue,
        "soon_count": soon,
        "linked_file_count": linked,
        "input_file_count": file_count,
        "budget_total": round(budget, 2),
    }


def _summary_columns():
    return (
        ContentAsset.id, ContentAsset.client_name, ContentAsset.campaign_name,
        ContentAsset.brand_id, ContentAsset.responsible_member_id,
        ContentAsset.responsible_member_ids, ContentAsset.period_start,
        ContentAsset.period_end, ContentAsset.status, ContentAsset.input_files,
        ContentAsset.kols, ContentAsset.assigned_user_ids, ContentAsset.budget_show,
    )


@router.get("", response_model=ContentAssetList)
def list_assets(
    search: str | None = None,
    owner_email: str | None = None,
    status: str | None = None,
    brand_id: int | None = None,
    skip: int = 0,
    limit: int = Query(200, le=500),
    user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    stmt = _filter_asset_query(select(ContentAsset), search, owner_email, status, brand_id)
    stmt = stmt.order_by(ContentAsset.updated_at.desc())
    if user.role == "admin":
        count_stmt = _filter_asset_query(
            select(func.count(ContentAsset.id)), search, owner_email, status, brand_id
        )
        total = db.scalar(count_stmt) or 0
        items = db.execute(stmt.offset(skip).limit(limit)).scalars().all()
    else:
        # JSON membership differs across SQLite/Postgres, so retain portable
        # filtering for scoped users while selecting only their final page.
        rows = db.execute(stmt).scalars().all()
        rows = [a for a in rows if user.id in (a.assigned_user_ids or [])]
        total = len(rows)
        items = rows[skip: skip + limit]
    for it in items:
        _redact_budget(it, user)
    return {"total": total, "items": items}


@router.get("/dashboard")
def campaign_dashboard(
    user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Aggregate dashboard data server-side and return only the top attention rows."""
    rows = db.execute(select(*_summary_columns())).all()
    if user.role != "admin":
        rows = [row for row in rows if user.id in (row.assigned_user_ids or [])]
    summaries = [_summary_row(row, user) for row in rows]
    attention = [
        {
            "id": item["id"], "campaign_name": item["campaign_name"],
            "client_name": item["client_name"], "overdue": item["overdue_count"],
            "soon": item["soon_count"],
        }
        for item in summaries
        if item["status"] not in {"completed", "complete", "done", "success", "succeeded", "cancelled"}
        and (item["overdue_count"] or item["soon_count"])
    ]
    attention.sort(key=lambda item: (item["overdue"], item["soon"]), reverse=True)
    return {
        "active": sum(item["status"] == "active" for item in summaries),
        "pending": sum(item["pending_count"] for item in summaries),
        "overdue": sum(item["overdue_count"] for item in summaries),
        "budget_total": round(sum(item["budget_total"] for item in summaries), 2) if user.role == "admin" else 0,
        "attention": attention[:8],
    }


@router.get("/summaries", response_model=ContentAssetSummaryList)
def list_asset_summaries(
    search: str | None = None,
    owner_email: str | None = None,
    status: str | None = None,
    brand_id: int | None = None,
    skip: int = 0,
    limit: int = Query(200, le=500),
    user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Small list payload for cards, dashboards, search, and export pickers."""
    stmt = _filter_asset_query(select(*_summary_columns()), search, owner_email, status, brand_id)
    stmt = stmt.order_by(ContentAsset.updated_at.desc())
    if user.role == "admin":
        count_stmt = _filter_asset_query(
            select(func.count(ContentAsset.id)), search, owner_email, status, brand_id
        )
        total = db.scalar(count_stmt) or 0
        rows = db.execute(stmt.offset(skip).limit(limit)).all()
    else:
        visible = [row for row in db.execute(stmt).all() if user.id in (row.assigned_user_ids or [])]
        total = len(visible)
        rows = visible[skip: skip + limit]
    return {"total": total, "items": [_summary_row(row, user) for row in rows]}


@router.get("/{asset_id}", response_model=ContentAssetOut)
def get_asset(asset_id: int, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(ContentAsset, asset_id)
    if not obj or not _can_read(user, obj):
        raise HTTPException(404, "Content asset not found")   # hide existence from unauthorised users
    _hydrate_kol_names(obj, db)
    _redact_budget(obj, user)
    return obj


@router.post("", response_model=ContentAssetOut, status_code=201)
def create_asset(data: ContentAssetCreate, user: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    values = data.model_dump()
    if not values.get("input_files"):
        values["input_files"] = [dict(f) for f in DEFAULT_INPUT_FILES]
    # JSON list columns must never be None (response model requires lists).
    for k in ("kols", "sow_options", "influencer_ids", "assigned_user_ids", "responsible_member_ids"):
        if values.get(k) is None:
            values[k] = []
    if values.get("budget_show") is None:
        values["budget_show"] = {}
    obj = ContentAsset(**values)
    _enforce_company(db, obj)
    _sanitize_urls(obj)
    db.add(obj)
    db.flush()   # assign id without committing, so the audit row lands atomically
    _log(db, obj.id, user, "created", f"สร้างแคมเปญ: {obj.campaign_name}")
    db.commit()
    db.refresh(obj)
    return obj


@router.get("/{asset_id}/history")
def asset_history(asset_id: int, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Audit trail for one campaign (who changed what, when)."""
    obj = db.get(ContentAsset, asset_id)
    if not obj or not _can_read(user, obj):
        raise HTTPException(404, "Content asset not found")
    rows = db.query(models.ChangeLog).filter(
        models.ChangeLog.entity == "campaign", models.ChangeLog.asset_id == asset_id
    ).order_by(models.ChangeLog.created_at.desc()).limit(50).all()
    return [{"actor": r.actor, "actor_id": r.actor_id, "action": r.action,
             "summary": r.summary, "detail": r.detail, "at": r.created_at.isoformat() + "Z"} for r in rows]


@router.put("/{asset_id}", response_model=ContentAssetOut)
def update_asset(
    asset_id: int,
    data: ContentAssetUpdate,
    user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    obj = db.get(ContentAsset, asset_id)
    if not obj or not _can_read(user, obj):
        raise HTTPException(404, "Content asset not found")
    if not _can_edit(user, obj):
        raise HTTPException(403, "คุณมีสิทธิ์ดูแคมเปญนี้เท่านั้น (แก้ไขไม่ได้)")
    changes = data.model_dump(exclude_unset=True)
    # Optimistic concurrency: if the client sent the version it loaded and the row
    # has since moved on, reject rather than silently overwrite the other edit.
    expected = changes.pop("row_version", None)
    if expected is not None and (obj.row_version or 1) != expected:
        raise HTTPException(409, "แคมเปญถูกแก้ไขโดยผู้อื่นไปแล้ว — โปรดโหลดเวอร์ชันล่าสุดก่อนบันทึก")
    if user.role != "admin":
        changes = {k: v for k, v in changes.items() if k not in ADMIN_ONLY_FIELDS}
    # JSON list/dict columns must never be set to None (the response model requires
    # them) — an explicit null in the body coerces back to the empty container.
    for k in ("kols", "sow_options", "influencer_ids", "assigned_user_ids", "responsible_member_ids", "input_files"):
        if k in changes and changes[k] is None:
            changes[k] = []
    if "budget_show" in changes and changes["budget_show"] is None:
        changes["budget_show"] = {}
    old_status = obj.status
    for key, value in changes.items():
        setattr(obj, key, value)
    _enforce_company(db, obj)
    _sanitize_urls(obj)
    if changes:
        obj.row_version = (obj.row_version or 1) + 1
        detail = ({"status": {"from": old_status, "to": obj.status}}
                  if "status" in changes and old_status != obj.status else None)
        _log(db, obj.id, user, "updated", "แก้ไข: " + ", ".join(sorted(changes.keys())), detail=detail)
    db.commit()
    db.refresh(obj)
    return obj


@router.put("/{asset_id}/drive-links", response_model=ContentAssetOut)
def update_drive_links(
    asset_id: int,
    links: dict[str, str],
    user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Set the Drive folder URL per input-file key; marks linked when a URL is present."""
    obj = db.get(ContentAsset, asset_id)
    if not obj or not _can_read(user, obj):
        raise HTTPException(404, "Content asset not found")
    if not _can_edit(user, obj):
        raise HTTPException(403, "คุณมีสิทธิ์ดูแคมเปญนี้เท่านั้น (แก้ไขไม่ได้)")
    files = [dict(f) for f in (obj.input_files or [])]
    for f in files:
        if f.get("key") in links:
            url = _safe_url(links[f["key"]])
            f["drive_url"] = url
            f["linked"] = bool(url)
    obj.input_files = files
    _log(db, obj.id, user, "updated", "อัปเดตลิงก์ Drive")
    db.commit()
    db.refresh(obj)
    return obj


# ---- KOL Plan import from the "KOLs Confirmed" Excel format ----------------
# Per-platform link columns carry a real hyperlink (the cell text is just "Link"),
# so we read cell.hyperlink.target rather than the displayed value.
_KOL_HEADER_MATCH_THRESHOLD = 95.0
_KOL_PLATFORM_FIELDS = ("tiktok", "instagram", "facebook", "lemon8", "youtube")
_KOL_FIXED_COLS = {
    "month": 1, "kol_type": 2, "name": 3, "profile_link": 4, "followers": 5,
    "content_type": 6, "sow": 7, "product_focus": 8, "post_date": 9,
    "tiktok": 10, "instagram": 11, "facebook": 12, "lemon8": 13, "youtube": 14,
    "kol_price": 16, "gencode_boosting": 17, "cart_added": 18, "buy_asset": 19,
    "outside_shooting": 20, "condition": 21, "gencode": 22,
}
_KOL_IMPORT_HEADERS = {
    "month": ("month", "campaign month", "period"),
    "kol_type": ("kols type", "kol type", "creator type", "tier", "type"),
    "name": ("kol name", "kols name", "kols/channel", "channel", "creator name", "creator", "influencer name", "name"),
    "profile_link": (
        "channels link", "channel link", "kol channel link", "creator channel link",
        "channel url", "profile link", "kol profile", "creator profile", "profile url", "profile",
    ),
    "followers": ("followers", "follower", "no. followers", "number of followers", "fan"),
    "content_type": ("content type", "content format", "content"),
    "sow": ("sow", "scope of work", "deliverable", "deliverables"),
    "product_focus": ("product focus", "product", "focus product", "focus"),
    "post_date": ("post date", "posting date", "publish date", "date"),
    "tiktok": ("tiktok", "tik tok", "tiktok link", "tik tok link"),
    "instagram": ("instagram", "ig", "instagram link", "ig link"),
    "facebook": ("facebook", "fb", "facebook link", "fb link"),
    "lemon8": ("lemon8", "lemon 8", "lemon8 link", "lemon 8 link"),
    "youtube": ("youtube", "you tube", "youtube link", "yt", "yt link"),
    "kol_price": ("kol price", "kol price thb", "kols price", "kols price thb", "price", "price thb", "rate", "fee"),
    "gencode_boosting": (
        "gencode boosting", "gencode/boosting", "gen code boosting",
        "gencode boosting thb", "gencode boost", "gen code boost", "boosting",
        "boosting cost", "boosting cost thb",
    ),
    "cart_added": ("cart added", "cart added thb", "add to cart", "cart"),
    "buy_asset": ("buy asset", "buy asset thb", "asset buyout", "asset buy out", "buyout"),
    "outside_shooting": ("outside shooting", "outside shooting thb", "outside", "shooting"),
    "condition": ("condition", "conditions", "terms"),
    "gencode": ("gencode", "gen code", "code"),
}
MAX_KOL_IMPORT_ROWS = 5000


def _header_text(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _header_key(value) -> str:
    return re.sub(r"[\W_]+", "", _header_text(value), flags=re.UNICODE)


def _header_score(value, alias: str) -> float:
    text = _header_text(value)
    alias_text = _header_text(alias)
    if not text or not alias_text:
        return 0.0
    if _header_key(text) == _header_key(alias_text):
        return 100.0
    return float(max(
        fuzz.ratio(_header_key(text), _header_key(alias_text)),
        fuzz.token_sort_ratio(text, alias_text),
    ))


def _best_kol_header(value) -> tuple[str | None, float]:
    best_field, best_score = None, 0.0
    for field, aliases in _KOL_IMPORT_HEADERS.items():
        for alias in aliases:
            score = _header_score(value, alias)
            if score > best_score:
                best_field, best_score = field, score
    if best_score >= _KOL_HEADER_MATCH_THRESHOLD:
        return best_field, best_score
    return None, best_score


def _map_kol_headers(ws) -> tuple[int | None, dict[str, int]]:
    best_row, best_map, best_score = None, {}, 0.0
    for row in range(1, min(ws.max_row, 12) + 1):
        row_map: dict[str, int] = {}
        row_scores: dict[str, float] = {}
        for col in range(1, ws.max_column + 1):
            field, score = _best_kol_header(ws.cell(row=row, column=col).value)
            if not field:
                continue
            if score > row_scores.get(field, -1.0):
                row_map[field] = col
                row_scores[field] = score
        if "name" not in row_map:
            continue
        score = (len(row_map) * 1000) + sum(row_scores.values())
        if score > best_score:
            best_row, best_map, best_score = row, row_map, score
    return best_row, best_map


def _parse_confirmed_kols(raw: bytes, max_rows: int = MAX_KOL_IMPORT_ROWS) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = next((wb[s] for s in wb.sheetnames if "confirm" in s.lower() or "comfirm" in s.lower()), wb.active)
    header_row, header_cols = _map_kol_headers(ws)
    cols = header_cols or _KOL_FIXED_COLS
    first_data_row = (header_row + 1) if header_row else 3

    def val(field, r):
        col = cols.get(field)
        if not col:
            return None
        v = ws.cell(row=r, column=col).value
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d")
        return v

    def link(field, r):
        col = cols.get(field)
        if not col:
            return None
        c = ws.cell(row=r, column=col)
        return c.hyperlink.target if c.hyperlink else None

    def num(field, r):
        v = val(field, r)
        try:
            return float(v) if v not in (None, "-", "") else 0
        except (TypeError, ValueError):
            return 0

    rows: list[dict] = []
    cur_month = ""
    for r in range(first_data_row, ws.max_row + 1):
        name = val("name", r)
        if not name or str(name).strip() in ("", "-"):
            continue
        m = val("month", r)
        if m:
            cur_month = str(m).strip()
        links = {}
        for plat in _KOL_PLATFORM_FIELDS:
            raw_link = link(plat, r) or val(plat, r)
            safe = _safe_url(raw_link)
            if safe:
                links[plat] = safe
        sow_val = val("sow", r)
        prof = link("profile_link", r) or (val("profile_link", r) if val("profile_link", r) not in ("-", None) else "")
        rows.append({
            "month": cur_month, "kol_type": (val("kol_type", r) or ""), "name": str(name).strip(),
            "profile_link": _safe_url(prof), "followers": (val("followers", r) or ""),
            "content_type": (val("content_type", r) or ""),
            "sow": [sow_val] if sow_val and sow_val != "-" else [],
            "product_focus": (val("product_focus", r) or ""), "post_date": (val("post_date", r) or ""),
            "links": links,
            "kol_price": num("kol_price", r), "gencode_boosting": num("gencode_boosting", r), "cart_added": num("cart_added", r),
            "buy_asset": num("buy_asset", r), "outside_shooting": num("outside_shooting", r),
            "condition": (val("condition", r) or ""), "gencode": (val("gencode", r) or ""),
        })
        if len(rows) > max_rows:
            raise ValueError(f"KOL import has more than {max_rows} rows")
    return rows


@router.post("/{asset_id}/import-kols")
async def import_kols(
    asset_id: int,
    file: UploadFile = File(...),
    user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Parse an uploaded 'KOLs Confirmed' .xlsx and RETURN the KOL rows (preview).
    The client appends/saves them — this endpoint does not mutate the campaign."""
    obj = db.get(ContentAsset, asset_id)
    if not obj or not _can_read(user, obj):
        raise HTTPException(404, "Content asset not found")
    if not _can_edit(user, obj):
        raise HTTPException(403, "คุณมีสิทธิ์ดูแคมเปญนี้เท่านั้น (นำเข้าไม่ได้)")
    raw = await file.read()
    ext = validate_spreadsheet_upload(
        file.filename, raw,
        allowed_extensions={".xlsx"},
        max_mb=settings.max_upload_mb,
    )
    if ext != ".xlsx":
        raise HTTPException(400, "รองรับเฉพาะไฟล์ .xlsx")
    try:
        kols = _parse_confirmed_kols(raw)
    except Exception as e:  # noqa: BLE001 — surface a friendly parse error
        raise HTTPException(400, f"อ่านไฟล์ไม่สำเร็จ: {e}")
    return {"count": len(kols), "kols": kols}


@router.delete("/{asset_id}", status_code=204)
def delete_asset(asset_id: int, user: models.User = Depends(require_admin), db: Session = Depends(get_db)):
    obj = db.get(ContentAsset, asset_id)
    if not obj:
        raise HTTPException(404, "Content asset not found")
    _log(db, asset_id, user, "deleted", f"ลบแคมเปญ: {obj.campaign_name}")
    db.delete(obj)
    db.commit()


@router.get("/{asset_id}/export")
def export_asset(
    asset_id: int,
    format: str = Query("xlsx", pattern="^(xlsx|csv|json|pdf|png)$"),
    user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Download one campaign workspace (admin or assigned manager only)."""
    from .. import models as _m
    obj = db.get(ContentAsset, asset_id)
    if not obj or not _can_read(user, obj):
        raise HTTPException(404, "Content asset not found")
    if not _can_edit(user, obj):
        raise HTTPException(403, "คุณมีสิทธิ์ดูแคมเปญนี้เท่านั้น (export ไม่ได้)")

    _hydrate_kol_names(obj, db)
    names = {i.id: i.name for i in db.execute(select(_m.Influencer)).scalars().all()}
    def spreadsheet_value(value):
        # Prevent imported user text from becoming a formula when opened in a
        # spreadsheet program. JSON exports preserve the original value.
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    rows = []
    for k in (obj.kols or []):
        rate, gen, boost = _num(k.get("rate")), _num(k.get("gen_code_price")), _num(k.get("boosting_cost"))
        rows.append({
            "Month": k.get("month", ""),
            "Tier": k.get("tier", ""),
            "Type": k.get("kol_type", ""),
            "KOL": names.get(k.get("influencer_id"), k.get("name", "")),
            "SOW": ", ".join(k.get("sow", []) if isinstance(k.get("sow"), list) else []),
            "Product Focus": k.get("product_focus", ""),
            "Approved": k.get("client_approved", ""),
            "Post Date": k.get("post_date", ""),
            "Link": k.get("link", ""),
            "ค่าตัว": rate,
            "Gen Code": gen,
            "Boosting": boost,
            "Total": round(rate + gen + boost, 2),
            "Objective": k.get("objective", ""),
        })
    df = pd.DataFrame(rows).map(spreadsheet_value)
    stamp = datetime.now().strftime("%Y%m%d")
    safe = "".join(
        c for c in (obj.campaign_name or "")
        if c.isascii() and (c.isalnum() or c in " -_")
    )[:40].strip() or f"campaign-{obj.id}"
    fname = f"{safe}_{stamp}"

    history_rows = db.query(models.ChangeLog).filter(
        models.ChangeLog.entity == "campaign", models.ChangeLog.asset_id == asset_id
    ).order_by(models.ChangeLog.created_at.desc()).all()
    history = [{
        "actor": r.actor,
        "actor_id": r.actor_id,
        "action": r.action,
        "summary": r.summary,
        "detail": r.detail,
        "at": r.created_at.isoformat() + "Z",
    } for r in history_rows]

    campaign_data = ContentAssetOut.model_validate(obj).model_dump(mode="json")
    if format in {"pdf", "png"}:
        try:
            report = render_campaign_report(campaign_data, rows, format)
        except ValueError as exc:
            raise HTTPException(413, str(exc)) from exc
        return StreamingResponse(
            report,
            media_type="application/pdf" if format == "pdf" else "image/png",
            headers={"Content-Disposition": f'attachment; filename="{fname}.{format}"'},
        )

    if format == "json":
        data = json.dumps({
            "export_version": 1,
            "exported_at": datetime.now().isoformat(),
            "campaign": campaign_data,
            "activity_log": history,
        }, ensure_ascii=False, indent=2).encode("utf-8")
        return StreamingResponse(
            io.BytesIO(data), media_type="application/json",
            headers={"Content-Disposition": f'attachment; filename="{fname}.json"'},
        )

    if format == "csv":
        data = df.to_csv(index=False).encode("utf-8-sig")
        return StreamingResponse(io.BytesIO(data), media_type="text/csv",
                                 headers={"Content-Disposition": f'attachment; filename="{fname}.csv"'})

    def excel_rows(items: list[dict]) -> list[dict]:
        return [{
            key: spreadsheet_value(
                json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value
            )
            for key, value in row.items()
        } for row in items if isinstance(row, dict)]

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        summary_fields = (
            "id", "campaign_name", "client_name", "brand_id", "owner_name", "owner_email",
            "period_start", "period_end", "status", "tags", "description", "stakeholders",
            "drive_folder_url", "responsible_member_id", "responsible_member_ids",
            "assigned_user_ids", "budget_show", "created_at", "updated_at",
        )
        summary = []
        for key in summary_fields:
            value = campaign_data.get(key, "")
            if isinstance(value, (list, dict)):
                value = json.dumps(value, ensure_ascii=False)
            summary.append({"Field": key, "Value": spreadsheet_value(value)})
        pd.DataFrame(summary).to_excel(writer, index=False, sheet_name="Campaign Summary")
        df.to_excel(writer, index=False, sheet_name="KOL Plan")
        pd.DataFrame(excel_rows(obj.performance_results or [])).to_excel(writer, index=False, sheet_name="Performance")
        pd.DataFrame(excel_rows(obj.input_files or [])).to_excel(writer, index=False, sheet_name="Input Files")
        pd.DataFrame(excel_rows(history)).to_excel(writer, index=False, sheet_name="Activity Log")
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'},
    )
