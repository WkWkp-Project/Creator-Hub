"""Spreadsheet import endpoints.

Flow:
  1. POST /api/imports/preview  -> upload CSV/XLSX, get auto-matched column map
  2. POST /api/imports/commit   -> confirm map, bulk upsert into the database

Uploaded files are cached under a temp dir keyed by a UUID so the commit
step can re-read them without a second upload.
"""
import json
import os
import re
import tempfile
import uuid
from io import BytesIO

import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from .. import crud, models, schemas
from ..config import get_settings
from ..database import get_db
from ..deps import require_admin
from ..file_validation import validate_spreadsheet_upload
from ..services import column_matcher as cm
from ..services.tiers import tier_for_followers

router = APIRouter(prefix="/api/imports", tags=["imports"])
settings = get_settings()

CACHE_DIR = os.path.join(tempfile.gettempdir(), "creatorhub_imports")
os.makedirs(CACHE_DIR, exist_ok=True)
MAX_IMPORT_ROWS = 5000   # guard against runaway files (DoS / DB bloat)
MAX_IMPORT_COLUMNS = 200


def _read_dataframe(path: str, filename: str) -> pd.DataFrame:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".csv":
        return pd.read_csv(path, dtype=str, keep_default_na=False)
    return pd.read_excel(path, dtype=str, keep_default_na=False)


def _validate_import_shape(df: pd.DataFrame) -> None:
    if len(df) > MAX_IMPORT_ROWS:
        raise HTTPException(400, f"ไฟล์มี {len(df)} แถว เกินลิมิต {MAX_IMPORT_ROWS} แถวต่อครั้ง")
    if len(df.columns) > MAX_IMPORT_COLUMNS:
        raise HTTPException(400, f"ไฟล์มี {len(df.columns)} คอลัมน์ เกินลิมิต {MAX_IMPORT_COLUMNS} คอลัมน์ต่อครั้ง")


@router.get("/system-fields")
def system_fields():
    """Expose the field catalogue so the UI can build mapping dropdowns."""
    return [{"value": k, "label": v} for k, v in cm.SYSTEM_FIELDS.items()]


# A ready-to-fill template — headers chosen so the auto-matcher maps them 100%,
# plus two example rows. No DB / sensitive data, so it needs no auth.
TEMPLATE_ROWS: list[tuple[str, str, str]] = [
    ("ชื่อ *", "Nong Aom", "Tee Talk"),
    ("Handle", "@nongaom", "@teetalk"),
    ("Niche", "Beauty", "Tech"),
    ("Platform", "Instagram", "YouTube"),
    ("Followers", "125000", "1.2M"),
    ("Engagement Rate", "4.8%", "3.1%"),
    ("Tier", "Micro", "Mega"),
    ("ค่าตัว", "95000", "300000"),
    ("ค่าเจนโค้ด", "10000", "20000"),
    ("ค่าเมเนจฟี", "5000", "15000"),
    ("ค่าเอเจนฟี %", "15", "20"),
    ("Currency", "THB", "THB"),
    ("Age", "26", "31"),
    ("Location", "Bangkok", "Chiang Mai"),
    ("Bio", "บิวตี้ครีเอเตอร์สายแต่งหน้า", "รีวิวแกดเจ็ต/มือถือ"),
    ("Instagram Link", "https://instagram.com/nongaom", ""),
    ("TikTok Link", "https://tiktok.com/@nongaom", ""),
    ("YouTube Link", "", "https://youtube.com/@teetalk"),
    ("Facebook Link", "", ""),
    ("Verified", "yes", "no"),
    ("Notes", "ลูกค้า VIP", ""),
]
_TEMPLATE_HELP = [
    "วิธีใช้เทมเพลตนำเข้าอินฟลูเอนเซอร์",
    "",
    "1) กรอกข้อมูลในชีต \"Influencers\" — แถวที่ 2-3 เป็นตัวอย่าง (สีเทา) ให้ลบทิ้งก่อนใช้จริง",
    "2) เก็บแถวหัวตาราง (แถวที่ 1) ไว้ — ระบบใช้แมชคอลัมน์อัตโนมัติ",
    "3) บันทึกไฟล์ แล้วอัปโหลดผ่านปุ่ม Import Data ในหน้า Directory",
    "",
    "กฎการกรอก:",
    "• ชื่อ = จำเป็น · คอลัมน์อื่นเว้นว่างได้",
    "• ตัวเลข: ใส่ 1.2M / ฿95,000 / 5,000 / 4.8% ได้ (ระบบแปลงให้)",
    "• Tier: Nano / Micro / Mega — เว้นว่าง = คำนวณจากยอด followers อัตโนมัติ",
    "• ค่าเอเจนฟี = เปอร์เซ็นต์ (ใส่ 15 = 15%)",
    "• Verified: yes / no",
    "• ชื่อหรือ handle ซ้ำกับที่มีอยู่ = อัปเดต, ไม่ซ้ำ = เพิ่มใหม่",
    "• หัวตารางสะกดใกล้เคียงก็พอ ระบบแมชภาษาไทย/อังกฤษให้",
]


@router.get("/template")
def template():
    """Download a fill-and-go Excel template with the recommended columns.

    Sheet 1 holds headers + two grey example rows (delete before importing);
    Sheet 2 holds the instructions, so guidance text never lands as import data.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Influencers"
    hdr_fill = PatternFill("solid", fgColor="E1121C")
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    ex_font = Font(color="9AA0A6", italic=True)   # examples look obviously like samples
    thin = Side(style="thin", color="E5E5E8")
    for c, (header, ex1, ex2) in enumerate(TEMPLATE_ROWS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin, right=thin)
        e1 = ws.cell(row=2, column=c, value=ex1); e1.font = ex_font
        e2 = ws.cell(row=3, column=c, value=ex2); e2.font = ex_font
        ws.column_dimensions[cell.column_letter].width = max(12, min(30, len(header) + 6))
    ws.freeze_panes = "A2"

    # Instructions on a separate sheet (the importer only reads the first sheet).
    ws2 = wb.create_sheet("วิธีใช้ (อ่านก่อน)")
    ws2.column_dimensions["A"].width = 90
    ws2.cell(row=1, column=1).font = Font(bold=True, size=13, color="E1121C")
    for i, line in enumerate(_TEMPLATE_HELP, start=1):
        ws2.cell(row=i, column=1, value=line)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="influencer_import_template.xlsx"'},
    )


@router.post("/preview", response_model=schemas.ImportPreview)
async def preview(file: UploadFile = File(...), _: models.User = Depends(require_admin)):
    raw = await file.read()
    ext = validate_spreadsheet_upload(
        file.filename, raw,
        allowed_extensions=settings.extensions,
        max_mb=settings.max_upload_mb,
    )

    upload_id = uuid.uuid4().hex
    stored = os.path.join(CACHE_DIR, f"{upload_id}{ext}")
    with open(stored, "wb") as f:
        f.write(raw)
    # remember the original filename for the commit step
    with open(os.path.join(CACHE_DIR, f"{upload_id}.meta"), "w") as f:
        json.dump({"filename": file.filename, "ext": ext}, f)

    try:
        df = _read_dataframe(stored, file.filename or "")
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"Could not parse file: {exc}") from exc
    _validate_import_shape(df)

    suggestions: list[schemas.ColumnSuggestion] = []
    for col in df.columns:
        outcome = cm.match_column(col)
        sample = [str(v) for v in df[col].head(3).tolist() if str(v).strip()]
        suggestions.append(
            schemas.ColumnSuggestion(
                file_column=str(col),
                system_field=outcome.system_field,
                confidence=outcome.confidence,
                status=outcome.status,
                sample=sample,
            )
        )

    mapped = sum(1 for s in suggestions if s.system_field)
    return schemas.ImportPreview(
        upload_id=upload_id,
        filename=file.filename or "upload",
        row_count=int(len(df)),
        detected_columns=[str(c) for c in df.columns],
        suggestions=suggestions,
        system_fields=list(cm.SYSTEM_FIELDS.keys()),
        mapped_count=mapped,
    )


@router.post("/commit", response_model=schemas.ImportResult)
def commit(payload: schemas.ImportCommit,
           actor: models.User = Depends(require_admin),
           db: Session = Depends(get_db)):
    # upload_id is a server-generated uuid4 hex — reject anything else so it can
    # never be used to traverse outside the cache dir.
    if not re.fullmatch(r"[0-9a-f]{32}", payload.upload_id or ""):
        raise HTTPException(400, "Invalid upload id.")
    meta_path = os.path.join(CACHE_DIR, f"{payload.upload_id}.meta")
    if not os.path.exists(meta_path):
        raise HTTPException(404, "Upload session expired or not found. Please re-upload.")
    meta = json.load(open(meta_path))
    stored = os.path.join(CACHE_DIR, f"{payload.upload_id}{meta['ext']}")

    df = _read_dataframe(stored, meta["filename"])
    _validate_import_shape(df)

    # file_column -> system_field (drop unmapped / null targets)
    mapping = {m.file_column: m.system_field for m in payload.mappings if m.system_field}
    if not mapping:
        raise HTTPException(400, "No columns mapped to system fields.")

    created = updated = skipped = 0
    errors: list[str] = []

    for idx, row in df.iterrows():
        try:
            record: dict = {}
            social: dict = {}
            for file_col, sys_field in mapping.items():
                if file_col not in df.columns:
                    continue
                raw = str(row[file_col]).strip()
                # Skip blank cells — never overwrite existing data with "" / 0,
                # and let model defaults apply on create.
                if not raw or raw.lower() in {"nan", "none", "-"}:
                    continue
                if sys_field in cm.LINK_FIELDS:
                    social[cm.LINK_FIELDS[sys_field]] = raw
                else:
                    record[sys_field] = cm.coerce(sys_field, row[file_col])
            if social:
                record["social_links"] = social

            name = record.get("name", "").strip()
            handle = record.get("handle", "").strip()
            if not name and not handle:
                skipped += 1
                continue

            # Auto-derive tier from followers when the file didn't provide one.
            if not record.get("tier") and "followers" in record:
                record["tier"] = tier_for_followers(record.get("followers"))

            existing = (
                crud.find_match(db, name, handle) if payload.update_existing else None
            )
            if existing:
                for k, v in record.items():
                    if k == "social_links" and existing.social_links:
                        merged = dict(existing.social_links)
                        merged.update(v)
                        existing.social_links = merged
                    else:
                        setattr(existing, k, v)
                updated += 1
            else:
                db.add(models.Influencer(**record))
                created += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"Row {int(idx) + 2}: {exc}")

    from .. import audit
    audit.record(db, entity="import", user=actor, action="commit",
                 summary=f"นำเข้าอินฟลู: created {created}, updated {updated}, skipped {skipped}",
                 detail={"created": created, "updated": updated, "skipped": skipped, "errors": len(errors)})
    db.commit()

    # cleanup cache
    for p in (stored, meta_path):
        try:
            os.remove(p)
        except OSError:
            pass

    return schemas.ImportResult(
        created=created, updated=updated, skipped=skipped, errors=errors[:25]
    )
